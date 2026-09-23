"""Generate a per-run Excel workbook of the pytest results and every API call
made during the run, via openpyxl. Styling mirrors the sibling UI-automation
reporter (Automation Saathi/reporters/excel_reporter.py) adapted for API rows.
"""
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Colour palette ───────────────────────────────────────────────────────
_DARK_BG = "0F1117"
_HEADER_BG = "1A1F2E"
_PASS_BG = "1C3A2A"
_FAIL_BG = "3A1C1C"
_XFAIL_BG = "2E1C3A"
_SKIP_BG = "3A2E1C"
_PASS_FG = "68D391"
_FAIL_FG = "FC8181"
_XFAIL_FG = "B794F6"
_SKIP_FG = "F6AD55"
_WHITE = "E2E8F0"
_MUTED = "718096"
_ACCENT = "667EEA"
_BORDER_COL = "2D3748"

_STATUS_STYLES = {
    "passed": (_PASS_FG, _PASS_BG),
    "PASS": (_PASS_FG, _PASS_BG),
    "failed": (_FAIL_FG, _FAIL_BG),
    "FAIL": (_FAIL_FG, _FAIL_BG),
    "xfailed": (_XFAIL_FG, _XFAIL_BG),
    "xpassed": (_XFAIL_FG, _XFAIL_BG),
    "skipped": (_SKIP_FG, _SKIP_BG),
}


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color=_WHITE, size=10) -> Font:
    return Font(bold=bold, color=color, size=size, name="Calibri")


def _thin_border() -> Border:
    side = Side(style="thin", color=_BORDER_COL)
    return Border(left=side, right=side, top=side, bottom=side)


def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def _write_header(ws, columns: list[tuple[str, int]]) -> None:
    ws.sheet_view.showGridLines = False
    for col_idx, (header, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 24
    for col_idx, (header, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _font(bold=True, color=_ACCENT)
        cell.fill = _fill(_HEADER_BG)
        cell.alignment = _center()
        cell.border = _thin_border()
    ws.freeze_panes = "A2"


# ── Summary sheet ─────────────────────────────────────────────────────────

def _write_summary(ws, summary: dict) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 50

    rows = [
        ("Environment", summary.get("environment", "")),
        ("Base URL", summary.get("base_url", "")),
        ("Mutations allowed", summary.get("allow_mutations", "")),
        ("Run started", summary.get("started_at", "")),
        ("Run finished", summary.get("finished_at", "")),
        ("Duration (s)", summary.get("duration_s", "")),
        ("", ""),
        ("Total tests", summary.get("total_tests", 0)),
        ("Passed", summary.get("passed", 0)),
        ("Failed", summary.get("failed", 0)),
        ("Known bugs (xfailed)", summary.get("xfailed", 0)),
        ("Skipped", summary.get("skipped", 0)),
        ("Pass rate", summary.get("pass_rate", "")),
        ("", ""),
        ("Total API calls", summary.get("total_calls", 0)),
        ("API call failures", summary.get("call_failures", 0)),
    ]
    ws.cell(row=1, column=1, value="Saathi API Automation - Run Summary").font = _font(bold=True, color=_ACCENT, size=13)
    ws.cell(row=1, column=1).fill = _fill(_HEADER_BG)
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 26

    for i, (label, value) in enumerate(rows, start=2):
        lcell = ws.cell(row=i, column=1, value=label)
        vcell = ws.cell(row=i, column=2, value=value)
        row_bg = _DARK_BG if i % 2 == 0 else _HEADER_BG
        for cell, align, font in ((lcell, _left(), _font(bold=True, color=_MUTED)), (vcell, _left(), _font(color=_WHITE))):
            cell.fill = _fill(row_bg)
            cell.alignment = align
            cell.border = _thin_border()
            cell.font = font


# ── API Calls sheet ───────────────────────────────────────────────────────

_CALL_COLUMNS = [
    ("#", 5),
    ("Test", 40),
    ("Service", 16),
    ("API / Operation", 32),
    ("Kind", 10),
    ("Method", 8),
    ("Path", 60),
    ("HTTP Status", 11),
    ("GraphQL Errors", 40),
    ("Latency (ms)", 12),
    ("Result", 9),
]
_CALL_COL_RESULT = 11


def _write_calls(ws, exchanges: list[dict]) -> None:
    _write_header(ws, _CALL_COLUMNS)
    for row_idx, ex in enumerate(exchanges, start=2):
        fg, bg = _STATUS_STYLES.get(ex["result"], (_WHITE, _DARK_BG))
        row_bg = _DARK_BG if row_idx % 2 == 0 else _HEADER_BG
        values = [
            row_idx - 1,
            ex["test"],
            ex["service"],
            ex["api"],
            ex["kind"],
            ex["method"],
            ex["path"],
            ex["status"],
            ex["graphql_errors"],
            round(ex["duration_ms"], 1),
            ex["result"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = _thin_border()
            if col_idx == _CALL_COL_RESULT:
                cell.font = _font(bold=True, color=fg)
                cell.fill = _fill(bg)
                cell.alignment = _center()
            elif col_idx in (1, 8, 10):
                cell.font = _font(color=_MUTED)
                cell.fill = _fill(row_bg)
                cell.alignment = _center()
            else:
                cell.font = _font(color=_FAIL_FG if col_idx == 9 and value else _WHITE)
                cell.fill = _fill(row_bg)
                cell.alignment = _left()


# ── Tests sheet ───────────────────────────────────────────────────────────

_TEST_COLUMNS = [
    ("#", 5),
    ("Test", 55),
    ("Status", 10),
    ("Duration (s)", 12),
    ("Error", 70),
]


def _write_tests(ws, test_results: list[dict]) -> None:
    _write_header(ws, _TEST_COLUMNS)
    for row_idx, r in enumerate(test_results, start=2):
        status = r["status"]
        fg, bg = _STATUS_STYLES.get(status, (_WHITE, _DARK_BG))
        row_bg = _DARK_BG if row_idx % 2 == 0 else _HEADER_BG
        error_text = (r.get("error") or "").split("\n")[0][:200]
        values = [row_idx - 1, r["name"], status.upper(), round(r.get("duration", 0) or 0, 2), error_text]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = _thin_border()
            if col_idx == 3:
                cell.font = _font(bold=True, color=fg)
                cell.fill = _fill(bg)
                cell.alignment = _center()
            elif col_idx in (1, 4):
                cell.font = _font(color=_MUTED)
                cell.fill = _fill(row_bg)
                cell.alignment = _center()
            else:
                cell.font = _font(color=_FAIL_FG if col_idx == 5 and error_text else _WHITE)
                cell.fill = _fill(row_bg)
                cell.alignment = _left()


# ── Public entry point ────────────────────────────────────────────────────

def generate_excel_report(
    summary: dict[str, Any],
    test_results: list[dict],
    exchanges: list[dict],
    output_path: str,
) -> str:
    """Generate the run's Excel workbook (Summary, API Calls, Tests sheets)
    and return the resolved output path."""
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    _write_summary(ws_summary, summary)

    ws_calls = wb.create_sheet("API Calls")
    _write_calls(ws_calls, exchanges)

    ws_tests = wb.create_sheet("Tests")
    _write_tests(ws_tests, test_results)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return str(Path(output_path).resolve())
